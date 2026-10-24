#!/usr/bin/env python3
"""Build the military drone events bundle for Skynet Research's
Resources tab from Draco's register workbook.

Input : Military_Drone_Events_Register_2026-10-09__1.xlsx
        (sheet "Events"; 7 title rows, header on row 8, data from row 9).
Output: assets/events.json + assets/events_manifest.json

The script fails loudly on any mismatch — a partial or misaligned
parse must never ship silently:
  - header row must match EXPECTED_HEADERS exactly
  - exactly 74 data rows, IDs E001..E074, unique and in order
  - every row carries a name, window, status, location, host,
    category, focus, evidence notes and a primary source URL
  - statuses must be within the register's own status set
  - every row's "Verified as of" is 2026-10-09

Re-run when Draco provides an updated register.
"""
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl is required: pip install openpyxl")

XLSX = Path.home() / ("workspace/user/files/"
                       "Military_Drone_Events_Register_2026-10-09__1.xlsx")
ROOT = Path(__file__).resolve().parent.parent
OUT_JSON = ROOT / "assets/events.json"
OUT_MANIFEST = ROOT / "assets/events_manifest.json"

EXPECTED_HEADERS = [
    "ID", "Event", "Published date / window", "Status", "Location",
    "Host / service", "Category", "Focus", "Participation / deadline",
    "Date / evidence notes", "Exact start", "Exact end",
    "Publisher / source type", "Primary source URL",
    "Supporting source URL", "Verified as of",
]
KNOWN_STATUSES = {
    "Completed", "Announced dates", "Announced window",
    "Past schedule only", "Tentative planning", "Date unconfirmed",
}
EXPECTED_COUNT = 74
VERIFIED = "2026-10-09"


def iso(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() or None


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Register workbook not found: {XLSX}")
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Events"]
    headers = [c.value for c in ws[8]]
    if headers != EXPECTED_HEADERS:
        raise SystemExit(f"Header mismatch:\n got={headers}\n "
                         f"expected={EXPECTED_HEADERS}")
    rows = []
    for r in ws.iter_rows(min_row=9, values_only=True):
        if all(v is None for v in r):
            continue
        rows.append(dict(zip(EXPECTED_HEADERS, r)))
    if len(rows) != EXPECTED_COUNT:
        raise SystemExit(f"Row count mismatch: {len(rows)} != "
                         f"{EXPECTED_COUNT}")
    events = []
    for i, row in enumerate(rows, 1):
        want_id = f"E{i:03d}"
        if row["ID"] != want_id:
            raise SystemExit(f"Row {i}: expected id {want_id}, "
                             f"got {row['ID']!r}")
        for field in ("Event", "Published date / window", "Status",
                      "Location", "Host / service", "Category", "Focus",
                      "Date / evidence notes", "Primary source URL"):
            if not row[field] or not str(row[field]).strip():
                raise SystemExit(f"{want_id}: missing {field}")
        if row["Status"] not in KNOWN_STATUSES:
            raise SystemExit(f"{want_id}: unknown status "
                             f"{row['Status']!r}")
        if iso(row["Verified as of"]) != VERIFIED:
            raise SystemExit(f"{want_id}: verified as of "
                             f"{iso(row['Verified as of'])}, not {VERIFIED}")
        events.append({
            "id": want_id,
            "name": str(row["Event"]).strip(),
            "window": str(row["Published date / window"]).strip(),
            "status": row["Status"],
            "location": str(row["Location"]).strip(),
            "host": str(row["Host / service"]).strip(),
            "category": row["Category"],
            "focus": str(row["Focus"]).strip(),
            "participation": (str(row["Participation / deadline"]).strip()
                              if row["Participation / deadline"] else None),
            "notes": str(row["Date / evidence notes"]).strip(),
            "exactStart": iso(row["Exact start"]),
            "exactEnd": iso(row["Exact end"]),
            "publisher": (str(row["Publisher / source type"]).strip()
                          if row["Publisher / source type"] else None),
            "primaryUrl": str(row["Primary source URL"]).strip(),
            "supportingUrl": (str(row["Supporting source URL"]).strip()
                              if row["Supporting source URL"] else None),
            "verifiedAsOf": VERIFIED,
        })
    starts = [e["exactStart"] for e in events if e["exactStart"]]
    payload = {
        "register": "Military Drone Events Register",
        "verifiedAsOf": VERIFIED,
        "note": ("Public-source inventory, verified 2026-10-09. Local, "
                 "classified, unannounced and older unit events may be "
                 "missing. Future dates remain subject to change; a "
                 "blank exact date means the full date is not verified."),
        "events": events,
    }
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False,
                                   separators=(",", ":")))
    manifest = {
        "register": payload["register"],
        "verifiedAsOf": VERIFIED,
        "entries": len(events),
        "byStatus": dict(Counter(e["status"] for e in events)),
        "byCategory": dict(Counter(e["category"] for e in events)),
        "exactStartRange": [min(starts), max(starts)],
        "withExactStart": len(starts),
        "withSupportingUrl": sum(1 for e in events if e["supportingUrl"]),
    }
    OUT_MANIFEST.write_text(json.dumps(manifest, indent=1,
                                       ensure_ascii=False))
    print(f"wrote {OUT_JSON} ({OUT_JSON.stat().st_size:,} bytes), "
          f"{len(events)} events")
    print(json.dumps(manifest, indent=1))


if __name__ == "__main__":
    main()
